#!/usr/bin/env python3
"""Build excel/reconciliatie-sales-finance-2026Q3.xlsx: the Q3 bridge between the weekly sales Excel and finance.

The workbook contains formulas, not pasted results: an analyst can open it, follow every step and
change the data. A "Controle" column compares each formula with the figure the SQL pipeline computed
(data/metrics.json) and shows OK when they match, so SQL and Excel check each other.
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "excel" / "reconciliatie-sales-finance-2026Q3.xlsx"
EUR = '€ #,##0.00;[Red]-€ #,##0.00'
BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="1F4E79")
WHITE = Font(bold=True, color="FFFFFF")
OKFILL = PatternFill("solid", fgColor="D9EAD3")
BADFILL = PatternFill("solid", fgColor="F4CCCC")


def header(ws, row, values, widths=None):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font, c.fill = WHITE, HEAD
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main() -> None:
    m = json.loads((ROOT / "data" / "metrics.json").read_text(encoding="utf-8"))
    rows = list(csv.reader(open(ROOT / "data" / "raw" / "legacy_sales_export_2026Q3.csv", encoding="utf-8")))[1:]
    fin = list(csv.reader(open(ROOT / "data" / "raw" / "finance_revenue_2026Q3.csv", encoding="utf-8")))[1:]
    wb = Workbook()

    # ------------------------------------------------------------------ Lees mij
    ws = wb.active
    ws.title = "Lees mij"
    lines = [
        ("Reconciliatie wekelijks sales-rapport ↔ finance, Q3 2026", True),
        ("Vennaert (fictief bedrijf) · illustratieve cijfers, geen echte bedrijfsdata", False),
        ("", False),
        ("Vraag van de sales manager: 'Waarom zegt ons Excel-rapport meer omzet dan finance?'", False),
        ("", False),
        ("Werkwijze", True),
        ("1. Tab Export = de ruwe export waarop het wekelijkse rapport gebouwd is (gesorteerd op datum, factuur, artikel).", False),
        ("2. Hulpkolommen G–J classificeren elke rij: soort document, dubbele rij, interne klant, netto bedrag met teken.", False),
        ("3. Tab Brug zet het rapporttotaal stap voor stap om naar het finance-cijfer. Elke stap is een SUMIFS op Export.", False),
        ("4. Kolom Controle vergelijkt elke stap met de SQL-pipeline (sql/05_reconciliation.sql). OK = beide tellen hetzelfde.", False),
        ("5. Tab Per verkoper = draaitabel met formules (netto per verkoper per maand). Echte draaitabel: Invoegen > Draaitabel op tabel tExport.", False),
        ("", False),
        ("Conclusie", True),
        ("Het rapport telt creditnota's op in plaats van ze af te trekken, bevat één dag twee keer en rekent de personeelsshop mee.", False),
        ("Na die drie correcties is het verschil met finance € 0,00. De nieuwe Power BI-rapportering gebruikt de finance-definitie (KPI-01).", False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold, size=14 if i == 1 else 11)
    ws.column_dimensions["A"].width = 120

    # ------------------------------------------------------------------ Export
    ex = wb.create_sheet("Export")
    header(ex, 1, ["Datum", "Factuur", "Klant", "Artikel", "Verkoper", "Bedrag", "Soort", "Dubbel", "Intern", "Netto"],
           [12, 11, 10, 11, 9, 14, 8, 8, 8, 14])
    for r, (d, inv, cust, sku, rep, amt) in enumerate(rows, 2):
        ex.cell(row=r, column=1, value=d)
        ex.cell(row=r, column=2, value=inv)
        ex.cell(row=r, column=3, value=cust)
        ex.cell(row=r, column=4, value=sku)
        ex.cell(row=r, column=5, value=rep)
        ex.cell(row=r, column=6, value=float(amt)).number_format = EUR
        ex.cell(row=r, column=7, value=f'=IF(LEFT(B{r},1)="C","Credit","Factuur")')
        # sorted on date, invoice, article: a repeated row sits right under its original
        ex.cell(row=r, column=8, value=0 if r == 2 else f"=IF(AND(B{r}=B{r-1},D{r}=D{r-1},A{r}=A{r-1},F{r}=F{r-1}),1,0)")
        ex.cell(row=r, column=9, value=f'=IF(C{r}="C99999",1,0)')
        c = ex.cell(row=r, column=10, value=f'=IF(OR(H{r}=1,I{r}=1),0,IF(G{r}="Credit",-F{r},F{r}))')
        c.number_format = EUR
    last = len(rows) + 1
    tab = Table(displayName="tExport", ref=f"A1:J{last}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ex.add_table(tab)
    ex.freeze_panes = "A2"

    # ------------------------------------------------------------------ Finance
    fi = wb.create_sheet("Finance")
    header(fi, 1, ["Periode", "Grootboek", "Omschrijving", "Bedrag"], [10, 11, 26, 16])
    for r, (period, gl, desc, amt) in enumerate(fin, 2):
        fi.cell(row=r, column=1, value=period)
        fi.cell(row=r, column=2, value=gl)
        fi.cell(row=r, column=3, value=desc)
        fi.cell(row=r, column=4, value=float(amt)).number_format = EUR
    fi.cell(row=len(fin) + 2, column=3, value="Totaal Q3").font = BOLD
    t = fi.cell(row=len(fin) + 2, column=4, value=f"=SUM(D2:D{len(fin) + 1})")
    t.number_format, t.font = EUR, BOLD

    # ------------------------------------------------------------------ Brug
    br = wb.create_sheet("Brug")
    header(br, 1, ["Stap", "Omschrijving", "Bedrag (Excel-formule)", "Rijen", "SQL-pipeline", "Controle"], [6, 50, 22, 9, 18, 11])
    E = f"Export!$F$2:$F${last}"
    G = f"Export!$G$2:$G${last}"
    H = f"Export!$H$2:$H${last}"
    C = f"Export!$C$2:$C${last}"
    steps = [
        (1, "Wekelijks Excel-rapport (som van de export)", f"=SUM({E})", f"=COUNTA(Export!$B$2:$B${last})", m["recon_legacy"]),
        (2, "Dubbel geëxporteerde dag eraf", f"=-SUMIFS({E},{H},1)", f"=SUM({H})", m["recon_dup_day"]),
        (3, "Creditnota's: opgeteld i.p.v. afgetrokken (2×)", f'=-2*SUMIFS({E},{G},"Credit",{H},0)', f'=COUNTIFS({G},"Credit",{H},0)', m["recon_credit"]),
        (4, "Personeelsshop (interne klant) eraf",
         f'=-(SUMIFS({E},{C},"C99999",{G},"Factuur",{H},0)-SUMIFS({E},{C},"C99999",{G},"Credit",{H},0))',
         f'=COUNTIFS({C},"C99999",{H},0)', m["recon_internal"]),
        (5, "Verklaard totaal", "=SUM(C2:C5)", None, m["recon_explained"]),
        (6, "Finance (grootboek 700000)", f"=Finance!D{len(fin) + 2}", None, m["recon_finance"]),
        (7, "Onverklaard verschil", "=ROUND(C6-C7,2)", None, m["recon_unexplained"]),
        (8, "Controle: som van kolom Netto op tab Export", f"=SUM(Export!$J$2:$J${last})", None, m["recon_finance"]),
    ]
    for i, (step, label, formula, count, sql_value) in enumerate(steps, 2):
        br.cell(row=i, column=1, value=step)
        br.cell(row=i, column=2, value=label).font = BOLD if step in (5, 7) else Font()
        br.cell(row=i, column=3, value=formula).number_format = EUR
        if count:
            br.cell(row=i, column=4, value=count)
        br.cell(row=i, column=5, value=sql_value).number_format = EUR
        br.cell(row=i, column=6, value=f'=IF(ROUND(C{i}-E{i},2)=0,"OK","VERSCHIL")').alignment = Alignment(horizontal="center")
    br.conditional_formatting.add("F2:F9", CellIsRule(operator="equal", formula=['"OK"'], fill=OKFILL))
    br.conditional_formatting.add("F2:F9", CellIsRule(operator="equal", formula=['"VERSCHIL"'], fill=BADFILL))
    br.cell(row=11, column=2, value="Kolom E is overgenomen uit data/metrics.json (SQL). Kolom C rekent Excel zelf uit.").font = Font(italic=True)

    # ------------------------------------------------------------------ Per verkoper (formula pivot)
    pv = wb.create_sheet("Per verkoper")
    reps = sorted({r[4] for r in rows if r[4]}) + [""]
    months = ["2026-07", "2026-08", "2026-09"]
    header(pv, 1, ["Verkoper", "jul", "aug", "sep", "Q3"], [18, 16, 16, 16, 16])
    for i, rep in enumerate(reps, 2):
        pv.cell(row=i, column=1, value=rep or "(leeg = niet toegewezen)")
        for j, month in enumerate(months, 2):
            c = pv.cell(row=i, column=j,
                        value=f'=SUMIFS(Export!$J$2:$J${last},Export!$E$2:$E${last},"{rep}",Export!$A$2:$A${last},"{month}*")' if rep else
                        f'=SUMIFS(Export!$J$2:$J${last},Export!$E$2:$E${last},"",Export!$A$2:$A${last},"{month}*")')
            c.number_format = EUR
        pv.cell(row=i, column=5, value=f"=SUM(B{i}:D{i})").number_format = EUR
    tot = len(reps) + 2
    pv.cell(row=tot, column=1, value="Totaal").font = BOLD
    for j in range(2, 6):
        col = get_column_letter(j)
        c = pv.cell(row=tot, column=j, value=f"=SUM({col}2:{col}{tot - 1})")
        c.number_format, c.font = EUR, BOLD
    pv.cell(row=tot + 2, column=1,
            value="Datums staan als tekst (JJJJ-MM-DD), daarom filtert \"2026-07*\" op juli. Lege verkoper: in het datawarehouse krijgt die lijn de verkoper op factuurdatum (regel ERP-01).").font = Font(italic=True)
    pv.freeze_panes = "B2"

    wb.save(OUT)
    print(f"written {OUT.relative_to(ROOT)} ({len(rows)} export rows)")


if __name__ == "__main__":
    main()
